from pathlib import Path
from datetime import datetime, date
import openpyxl, json
from sqlalchemy import select
from ..config import BASE_DIR
from ..models import *
from ..core.security import next_id, audit

EXCEL_DIR = BASE_DIR / 'excel'
SHEET_MAP = {
 'Servants':Servant,'Services':Service,'Service_Years':ServiceYear,'Users':User,'Roles':Role,'Permissions':Permission,'Role_Permissions':RolePermission,'User_Service_Scopes':UserServiceScope,
 'Service_Assignments':ServiceAssignment,'Service_Transfers':ServiceTransfer,'Meetings':Meeting,'Events':Event,'Meeting_Participants':Participant,
 'Servant_Meeting_Evaluations':Evaluation,'Excuses':Excuse,'Openings':Opening,'Followup_Records':FollowupRecord,
}
FILE_MAP={
 '01_CORE.xlsx':['Servants','Services','Service_Years','Users','Roles','Permissions','Role_Permissions','User_Service_Scopes'],
 '02_SERVICE_OPERATIONS.xlsx':['Service_Assignments','Service_Transfers','Meetings','Events','Meeting_Participants'],
 '03_EVALUATIONS.xlsx':['Servant_Meeting_Evaluations'],
 '04_EXCUSES.xlsx':['Excuses'],
 '05_OPENINGS.xlsx':['Openings'],
 '06_FOLLOWUP.xlsx':['Followup_Records'],
}
PK_MAP={
 'Servants':'servant_id','Services':'service_id','Service_Years':'service_year_id','Users':'user_id','Roles':'role_id','Permissions':'permission_id','Role_Permissions':'role_id', 'User_Service_Scopes':'scope_id','Service_Assignments':'assignment_id','Service_Transfers':'transfer_id','Meetings':'meeting_id','Events':'event_id','Meeting_Participants':'participant_id','Servant_Meeting_Evaluations':'evaluation_id','Excuses':'excuse_id','Openings':'opening_id','Followup_Records':'followup_id'}
FIELD_ALIASES={}

def _clean(v):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    return v

def _identity_from_row(model, vals, pk):
    """Return a SQLAlchemy identity value for simple/composite primary keys."""
    mapper = model.__mapper__
    keys = [col.key for col in mapper.primary_key] if mapper is not None else [pk]
    if len(keys) == 1:
        value = vals.get(keys[0])
        return str(value) if value not in (None, '') else None
    if not all(vals.get(k) not in (None, '') for k in keys):
        return None
    return tuple(str(vals[k]) for k in keys)


def _identity_label(identity):
    if isinstance(identity, tuple):
        return ':'.join(identity)
    return identity


def detect(db, user):
    changes=[]
    for fname,sheets in FILE_MAP.items():
        path=EXCEL_DIR/fname
        if not path.exists():
            continue
        wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
        try:
            for sheet in sheets:
                # Some workbook sheets are intentionally reference/config sheets and
                # are not part of the editable application model. Ignore them safely.
                if sheet not in wb.sheetnames or sheet not in SHEET_MAP or sheet not in PK_MAP:
                    continue
                ws=wb[sheet]
                header_row=next(ws.iter_rows(min_row=1,max_row=1), None)
                if not header_row:
                    continue
                headers=[c.value for c in header_row]
                model=SHEET_MAP[sheet]; pk=PK_MAP[sheet]
                mapper=model.__mapper__
                pk_keys=[col.key for col in mapper.primary_key]
                for row_no,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
                    vals={headers[i]:_clean(row[i]) if i<len(row) else None for i in range(len(headers)) if headers[i]}
                    if not any(v not in (None,'') for v in vals.values()):
                        continue
                    identity=_identity_from_row(model, vals, pk)
                    if identity is None:
                        changes.append((fname,sheet,None,'INSERT',None,vals,{'reason':'Primary key will be generated during Apply','row':row_no}))
                        continue
                    obj=db.get(model, identity)
                    if not obj:
                        changes.append((fname,sheet,_identity_label(identity),'INSERT',None,vals,{'reason':'new row','row':row_no}))
                        continue
                    old={c:getattr(obj,c,None) for c in headers if hasattr(obj,c)}
                    old={k:_clean(v) for k,v in old.items()}
                    diff={k:{'old':old.get(k),'new':vals.get(k)} for k in headers if k not in pk_keys and old.get(k)!=vals.get(k) and vals.get(k) not in (None,'')}
                    if diff:
                        changes.append((fname,sheet,_identity_label(identity),'UPDATE',old,vals,diff))
        finally:
            wb.close()
    return changes